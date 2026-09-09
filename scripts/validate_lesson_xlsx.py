#!/usr/bin/env python3
"""Validate a medical lesson-plan XLSX without modifying it."""

from __future__ import annotations

import argparse
import math
import re
import sys
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


HEADERS = [
    "日期",
    "章节",
    "教学目标",
    "教学重点",
    "教学难点",
    "复习提问",
    "新课引入",
    "新课讲授",
    "课堂小结",
    "布置作业",
]
MAJOR_NUMERALS = ["一", "二", "三", "四"]
GENERIC_PHRASES = (
    "核心概念与主要规律",
    "重点内容的比较方法与安全要点",
    "相关知识在健康服务中的规范沟通原则",
    "作用机制、临床用途与风险提示联系起来",
    "药物学习中应从哪些方面",
    "遇到同类药物信息时",
    "面对相关情境",
    "核心概念与分类",
    "应用与安全管理",
    "综合判断与课堂落实",
    "课堂上可用简短情境检验",
    "遇到资料不足的情况",
    "在健康服务情境中坚持核对信息和安全沟通",
    "用“分类、机制、应用、风险”四栏整理",
    "查阅一种相关药品说明书",
)
LECTURE_META_PHRASES = (
    "围绕",
    "本课将",
    "本节将",
    "课堂上",
    "教学中",
    "学习时",
    "可从",
    "可以从",
    "应从",
    "帮助学生",
    "便于理解",
    "便于记忆",
    "作为推理起点",
    "的学习不能只记",
    "的比较题应明确比较对象",
    "的知识要落实到病例线索",
    "建立“机制、效应、应用、风险”",
    "不是孤立名词",
    "涉及的判断要保留边界",
    "先核实给药途径、时间和合并用药",
    "不能直接下结论",
    "从治疗目标反推关键机制",
)
HOMEWORK_FILLER_PHRASES = (
    "围绕",
    "完成相关题目",
    "完成对应习题",
    "整理本课知识",
    "标注观察重点",
    "依次梳理全部知识点",
    "摘录适应证、禁忌和不良反应",
)
GENERIC_HEADINGS = {"核心概念与分类", "应用与安全管理", "综合判断与课堂落实"}
SIMILARITY_THRESHOLDS = {
    2: ("教学目标", 0.78),
    3: ("教学重点", 0.80),
    4: ("教学难点", 0.80),
    5: ("复习提问", 0.78),
    6: ("新课引入", 0.78),
    7: ("新课讲授", 0.72),
    8: ("课堂小结", 0.78),
    9: ("布置作业", 0.78),
}


def visible_count(value: str) -> int:
    """Count visible characters, including punctuation and numbering."""
    return len(re.sub(r"\s+", "", value))


def numbered_lines(value: str, expected: int) -> bool:
    lines = value.split("\n")
    if len(lines) != expected or any(not line.strip() for line in lines):
        return False
    return all(re.fullmatch(rf"{index}、.+", line) for index, line in enumerate(lines, 1))


def normalized_text(value: str, chapter: str = "") -> str:
    if chapter:
        value = value.replace(chapter, "本课标题")
    return re.sub(r"[\s，。；：、？！,.!?;:‘’“”《》（）()\-—]+", "", value)


def repeated_segments(value: str) -> list[str]:
    segments = []
    for segment in re.split(r"[。！？；\n]", value):
        segment = re.sub(r"^[一二三四\d]+、", "", segment.strip())
        normalized = normalized_text(segment)
        if len(normalized) >= 12:
            segments.append(normalized)
    counts = Counter(segments)
    return [segment for segment, count in counts.items() if count > 1]


def validate_lecture(value: str, chapter: str = "") -> list[str]:
    errors: list[str] = []
    count = visible_count(value)
    if not 800 <= count <= 1000:
        errors.append(f"新课讲授字数为 {count}，应为 800–1000")

    lines = value.split("\n")
    if not lines or not lines[0].startswith("一、"):
        errors.append("新课讲授必须以“一、”开始")
        return errors

    major_positions: list[int] = []
    major_labels: list[str] = []
    for index, line in enumerate(lines):
        match = re.fullmatch(r"([一二三四])、(.+)", line)
        if match:
            major_positions.append(index)
            major_labels.append(match.group(1))

    if not 3 <= len(major_positions) <= 4:
        errors.append("新课讲授必须有 3–4 个一级标题")
        return errors

    expected_labels = MAJOR_NUMERALS[: len(major_labels)]
    if major_labels != expected_labels:
        errors.append("新课讲授一级标题必须从“一、”开始连续编号")

    allowed_nonempty = re.compile(r"(?:[一二三四]、.+|\d+、.+)")
    for line in lines:
        if line and not allowed_nonempty.fullmatch(line):
            errors.append(f"新课讲授存在未按层级编号的行：{line[:24]}")
            break

    boundaries = major_positions + [len(lines)]
    total_items = 0
    for section_index, start in enumerate(major_positions):
        end = boundaries[section_index + 1]
        content_end = end
        if section_index < len(major_positions) - 1:
            if end < 2 or lines[end - 1] != "" or lines[end - 2] == "":
                errors.append("两个一级部分之间必须恰好留一个空行")
            content_end = end - 1
        section_lines = lines[start + 1 : content_end]
        if any(not line for line in section_lines):
            errors.append(f"第 {section_index + 1} 个一级部分内部不得留空行")
            continue
        total_items += len(section_lines)
        if not 3 <= len(section_lines) <= 4:
            errors.append(f"第 {section_index + 1} 个一级部分必须有 3–4 个要点")
            continue
        for item_number, line in enumerate(section_lines, 1):
            if not re.fullmatch(rf"{item_number}、.+", line):
                errors.append(
                    f"第 {section_index + 1} 个一级部分的要点必须从“1、”连续编号"
                )
                break
            point_count = visible_count(line)
            if not 35 <= point_count <= 110:
                errors.append(
                    f"第 {section_index + 1} 个一级部分第 {item_number} 点字数为 {point_count}，应为 35–110"
                )

        heading = re.sub(r"^[一二三四]、", "", lines[start])
        if heading in GENERIC_HEADINGS:
            errors.append(f"一级标题“{heading}”过于通用，必须改为本课具体知识标题")

    if total_items < 9:
        errors.append(f"新课讲授共有 {total_items} 个小点，至少需要 9 个")

    duplicates = repeated_segments(value)
    if duplicates:
        errors.append(f"新课讲授存在重复句或分句：{duplicates[0][:24]}")
    if chapter and chapter in value:
        errors.append("新课讲授不得复述完整章节标题")

    found_meta = [phrase for phrase in LECTURE_META_PHRASES if phrase in value]
    if found_meta:
        shown = "、".join(f"“{phrase}”" for phrase in found_meta[:5])
        errors.append(f"新课讲授含教学过程套话或可替换句式：{shown}")

    if lines[-1] == "":
        errors.append("新课讲授末尾不得有空行")
    return errors


def validate_row(values: list[object], excel_row: int) -> list[str]:
    errors: list[str] = []
    prefix = f"第 {excel_row} 行"
    if any(value is None or str(value).strip() == "" for value in values):
        errors.append(f"{prefix}存在空白必填单元格")
        return errors
    if any(not isinstance(value, str) for value in values):
        errors.append(f"{prefix}所有字段（尤其日期）必须保存为文本")
        return errors

    date, chapter, objective, focus, difficulty, review, intro, lecture, summary, homework = values
    fields = {
        "日期": date,
        "章节": chapter,
        "教学目标": objective,
        "教学重点": focus,
        "教学难点": difficulty,
        "复习提问": review,
        "新课引入": intro,
        "新课讲授": lecture,
        "课堂小结": summary,
        "布置作业": homework,
    }
    for name, value in fields.items():
        if "\r" in value:
            errors.append(f"{prefix}{name}含 CR/CRLF，必须只使用 LF 换行")

    no_break_fields = {
        "日期": date,
        "章节": chapter,
        "教学目标": objective,
        "教学重点": focus,
        "教学难点": difficulty,
        "新课引入": intro,
    }
    for name, value in no_break_fields.items():
        if "\n" in value:
            errors.append(f"{prefix}{name}不得换行")

    if not re.fullmatch(r"掌握.+，熟悉.+，了解.+。", objective) or objective.count("。") != 1:
        errors.append(f"{prefix}教学目标必须是“掌握…，熟悉…，了解…。”一个完整句子")
    if review and not numbered_lines(review, 2):
        errors.append(f"{prefix}复习提问必须是连续两行“1、”“2、”")

    intro_count = visible_count(intro)
    if not 100 <= intro_count <= 125:
        errors.append(f"{prefix}新课引入字数为 {intro_count}，应为 100–125")

    combined_authored = "\n".join([objective, focus, difficulty, review, intro, lecture, summary, homework])
    matched_generic = [phrase for phrase in GENERIC_PHRASES if phrase in combined_authored]
    if matched_generic:
        shown = "、".join(f"“{phrase}”" for phrase in matched_generic[:4])
        suffix = f"等 {len(matched_generic)} 处" if len(matched_generic) > 4 else ""
        errors.append(f"{prefix}存在通用填充语：{shown}{suffix}")

    for name, value in {
        "教学目标": objective,
        "教学重点": focus,
        "教学难点": difficulty,
        "课堂小结": summary,
    }.items():
        if chapter in value:
            errors.append(f"{prefix}{name}不得复述完整章节标题")

    errors.extend(f"{prefix}{message}" for message in validate_lecture(lecture, chapter))

    if summary and not numbered_lines(summary, 3):
        errors.append(f"{prefix}课堂小结必须是连续三行“1、”“2、”“3、”")
    if homework and not numbered_lines(homework, 3):
        errors.append(f"{prefix}布置作业必须是连续三行“1、”“2、”“3、”")
    else:
        homework_lines = homework.split("\n")
        first = homework_lines[0]
        if not any(keyword in first for keyword in ("教材", "书后", "课本", "习题")):
            errors.append(f"{prefix}布置作业第 1 项必须明确教材习题巩固任务")
        for index, line in enumerate(homework_lines, 1):
            length = visible_count(line)
            if not 18 <= length <= 90:
                errors.append(f"{prefix}布置作业第 {index} 项字数为 {length}，一般应为 18–90")
        found_homework_filler = [phrase for phrase in HOMEWORK_FILLER_PHRASES if phrase in homework]
        if found_homework_filler:
            shown = "、".join(f"“{phrase}”" for phrase in found_homework_filler[:4])
            errors.append(f"{prefix}布置作业含僵化或笼统表达：{shown}")
        if chapter in homework:
            errors.append(f"{prefix}布置作业不得复述完整章节标题")

    return errors


def populated_rows(sheet) -> Iterable[int]:
    for row_index in range(2, sheet.max_row + 1):
        if any(sheet.cell(row_index, column).value not in (None, "") for column in range(1, 11)):
            yield row_index


def validate_cross_row_quality(records: list[tuple[int, list[str]]]) -> list[str]:
    errors: list[str] = []
    for column_index, (field_name, threshold) in SIMILARITY_THRESHOLDS.items():
        exact_groups: dict[str, list[int]] = defaultdict(list)
        for excel_row, values in records:
            exact_groups[values[column_index]].append(excel_row)
        for matching_rows in exact_groups.values():
            if len(matching_rows) > 1:
                shown = "、".join(map(str, matching_rows[:10]))
                suffix = "等" if len(matching_rows) > 10 else ""
                errors.append(f"{field_name}在第 {shown}{suffix} 行完全重复")

        near_pairs: list[tuple[int, int, float]] = []
        for left_index in range(len(records)):
            left_row, left_values = records[left_index]
            left_text = normalized_text(left_values[column_index], left_values[1])
            for right_index in range(left_index + 1, len(records)):
                right_row, right_values = records[right_index]
                if left_values[column_index] == right_values[column_index]:
                    continue
                right_text = normalized_text(right_values[column_index], right_values[1])
                ratio = SequenceMatcher(None, left_text, right_text).ratio()
                if ratio >= threshold:
                    near_pairs.append((left_row, right_row, ratio))
        if near_pairs:
            examples = "；".join(
                f"第 {left}/{right} 行 {ratio:.0%}" for left, right, ratio in near_pairs[:5]
            )
            errors.append(
                f"{field_name}存在 {len(near_pairs)} 组高相似课次（阈值 {threshold:.0%}）：{examples}"
            )

    errors.extend(validate_repeated_long_fragments(records, 7, "新课讲授"))
    errors.extend(validate_repeated_long_fragments(records, 9, "布置作业"))
    errors.extend(validate_homework_task_mix(records))
    return errors


def validate_repeated_long_fragments(
    records: list[tuple[int, list[str]]], column_index: int, field_name: str
) -> list[str]:
    """Find sentence-frame fragments reused across many lessons despite noun substitution."""
    if len(records) < 3:
        return []
    fragment_rows: dict[str, set[int]] = defaultdict(set)
    width = 16
    for excel_row, values in records:
        text = normalized_text(values[column_index], values[1])
        for start in range(max(0, len(text) - width + 1)):
            fragment = text[start : start + width]
            if len(fragment) == width:
                fragment_rows[fragment].add(excel_row)
    minimum = max(3, math.ceil(len(records) * 0.30))
    repeated = [
        (fragment, sorted(rows))
        for fragment, rows in fragment_rows.items()
        if len(rows) >= minimum
    ]
    if not repeated:
        return []
    repeated.sort(key=lambda item: (-len(item[1]), item[0]))
    fragment, rows = repeated[0]
    shown = "、".join(map(str, rows[:8])) + ("等" if len(rows) > 8 else "")
    return [f"{field_name}存在跨课反复套用的长句片段“{fragment}”（第 {shown} 行）"]


def homework_task_signature(value: str) -> tuple[str, ...]:
    """Classify task forms so a whole course cannot reuse one homework recipe."""
    categories = {
        "比较表卡": ("比较表", "比较卡", "对照表"),
        "说明书摘录": ("说明书", "摘录"),
        "机制图": ("机制图", "流程图", "示意图"),
        "病例判断": ("病例", "案例", "判断依据"),
        "处方审核": ("处方", "医嘱", "审核"),
        "监测方案": ("监测", "观察记录", "随访"),
        "计算判读": ("计算", "曲线", "判读"),
        "概念网络": ("概念图", "思维导图", "知识网络"),
        "错题订正": ("错题", "订正", "纠错"),
        "口头讲解": ("口述", "讲解", "录音"),
        "资料查证": ("检索", "查证", "指南", "药典"),
    }
    found = [name for name, keywords in categories.items() if any(k in value for k in keywords)]
    return tuple(sorted(found)) or ("未识别任务型",)


def validate_homework_task_mix(records: list[tuple[int, list[str]]]) -> list[str]:
    if len(records) < 4:
        return []
    groups: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for excel_row, values in records:
        groups[homework_task_signature(values[9])].append(excel_row)
    limit = max(3, math.ceil(len(records) * 0.35))
    repeated = [(signature, rows) for signature, rows in groups.items() if len(rows) > limit]
    if not repeated:
        return []
    repeated.sort(key=lambda item: -len(item[1]))
    signature, rows = repeated[0]
    shown = "、".join(map(str, rows[:8])) + ("等" if len(rows) > 8 else "")
    return [
        f"布置作业任务组合“{'＋'.join(signature)}”在 {len(rows)} 课反复出现（第 {shown} 行），应按知识目标轮换任务形式"
    ]


def validate_loaded_workbook(workbook, expected_rows: int | None) -> list[str]:
    errors: list[str] = []
    if workbook.sheetnames != ["教案数据"]:
        errors.append(f"工作簿必须且只能包含“教案数据”，当前为：{workbook.sheetnames}")
        if "教案数据" not in workbook.sheetnames:
            return errors
    sheet = workbook["教案数据"]

    headers = [sheet.cell(1, column).value for column in range(1, 11)]
    if headers != HEADERS:
        errors.append(f"表头或顺序不正确：{headers}")
    if sheet.max_column != 10:
        errors.append(f"工作表必须恰好 10 列，当前检测到 {sheet.max_column} 列")

    row_numbers = list(populated_rows(sheet))
    if row_numbers and row_numbers != list(range(2, row_numbers[-1] + 1)):
        errors.append("数据区域中存在空白记录行")
    if expected_rows is not None and len(row_numbers) != expected_rows:
        errors.append(f"记录数为 {len(row_numbers)}，授课计划要求 {expected_rows}")

    records: list[tuple[int, list[str]]] = []
    for row_index in row_numbers:
        cells = [sheet.cell(row_index, column) for column in range(1, 11)]
        if any(cell.data_type == "f" for cell in cells):
            errors.append(f"第 {row_index} 行不得包含公式")
        for cell in cells:
            if cell.alignment.wrap_text is not True:
                errors.append(f"{cell.coordinate} 未启用自动换行")
            if cell.alignment.vertical != "top":
                errors.append(f"{cell.coordinate} 未设置为顶端对齐")
        values = [cell.value for cell in cells]
        errors.extend(validate_row(values, row_index))
        if all(isinstance(value, str) for value in values):
            records.append((row_index, values))

    errors.extend(validate_cross_row_quality(records))

    return errors


def validate_workbook(path: Path, expected_rows: int | None) -> list[str]:
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        return [f"无法读取工作簿：{exc}"]
    try:
        return validate_loaded_workbook(workbook, expected_rows)
    finally:
        workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="校验教案邮件合并 Excel")
    parser.add_argument("workbook", type=Path, help="待校验的 .xlsx 文件")
    parser.add_argument("--expected-rows", type=int, default=None, help="授课计划记录数")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    errors = validate_workbook(args.workbook, args.expected_rows)
    if errors:
        print(f"校验失败，共 {len(errors)} 项：")
        for error in errors:
            print(f"- {error}")
        return 1
    print(f"校验通过：{args.workbook}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
