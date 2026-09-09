#!/usr/bin/env python3
"""Regression tests for the lesson-plan workbook validator."""

from __future__ import annotations

import unittest

from openpyxl import Workbook
from openpyxl.styles import Alignment

from validate_lesson_xlsx import HEADERS, validate_loaded_workbook


def valid_lecture() -> str:
    characters = ["甲", "乙", "丙", "丁", "戊", "己", "庚", "辛", "壬"]
    sections = []
    for section_number, heading in enumerate(["受体作用", "体内过程", "用药辨析"]):
        points = []
        for point_number in range(1, 4):
            char = characters[section_number * 3 + point_number - 1]
            points.append(f"{point_number}、{char * 88}")
        sections.append(f"{'一二三'[section_number]}、{heading}\n" + "\n".join(points))
    return "\n\n".join(sections)


def valid_values() -> list[object]:
    return [
        "3月3日",
        "第一章：测试章节",
        "掌握受体激动与阻断的判定方法，熟悉量效关系曲线的读取步骤，了解治疗指数的安全意义。",
        "量效曲线中效价强度、效能与半数有效量的辨析。",
        "依据曲线位置和最大效应区分效价强度与效能。",
        "1、受体激动剂与拮抗剂对细胞效应有何不同？\n2、剂量增加时药物效应通常怎样变化？",
        "引" * 100,
        valid_lecture(),
        "1、总结概念。\n2、总结方法。\n3、总结应用。",
        "1、完成教材中量效曲线判读习题，并写出效价强度的判断依据。\n2、根据给定曲线判断三种药物的效能高低，逐项说明曲线证据。\n3、查证一种治疗窗较窄药物的监测要求，整理为六十字记录。",
    ]


def make_workbook(rows: list[list[object]] | None = None) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "教案数据"
    sheet.append(HEADERS)
    for values in rows or [valid_values()]:
        sheet.append(values)
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=10):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return workbook


class ValidatorTests(unittest.TestCase):
    def assert_has(self, errors: list[str], text: str) -> None:
        self.assertTrue(any(text in error for error in errors), errors)

    def test_valid_workbook_passes(self) -> None:
        self.assertEqual(validate_loaded_workbook(make_workbook(), 1), [])

    def test_wrong_header_is_detected(self) -> None:
        workbook = make_workbook()
        workbook.active["A1"] = "错误日期"
        self.assert_has(validate_loaded_workbook(workbook, 1), "表头或顺序不正确")

    def test_missing_row_is_detected(self) -> None:
        self.assert_has(validate_loaded_workbook(make_workbook(), 2), "记录数为 1")

    def test_empty_cell_is_detected(self) -> None:
        values = valid_values()
        values[6] = ""
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "空白必填单元格")

    def test_length_limits_are_detected(self) -> None:
        values = valid_values()
        values[6] = "短" * 99
        values[7] = values[7] + "长" * 200
        errors = validate_loaded_workbook(make_workbook([values]), 1)
        self.assert_has(errors, "新课引入字数")
        self.assert_has(errors, "新课讲授字数")

    def test_wrong_numbering_is_detected(self) -> None:
        values = valid_values()
        values[5] = "1、问题一？\n3、问题三？"
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "复习提问必须")

    def test_missing_line_break_is_detected(self) -> None:
        values = valid_values()
        values[8] = "1、总结概念。2、总结方法。3、总结应用。"
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "课堂小结必须")

    def test_numeric_date_is_detected(self) -> None:
        values = valid_values()
        values[0] = 46000
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "必须保存为文本")

    def test_missing_wrap_is_detected(self) -> None:
        workbook = make_workbook()
        workbook.active["A2"].alignment = Alignment(vertical="top")
        self.assert_has(validate_loaded_workbook(workbook, 1), "A2 未启用自动换行")

    def test_too_few_subpoints_are_detected(self) -> None:
        values = valid_values()
        values[7] = values[7].replace("\n3、" + "丙" * 88, "", 1)
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "必须有 3–4 个要点")

    def test_repeated_clause_is_detected(self) -> None:
        values = valid_values()
        repeated = "受体结合后产生可观察的组织效应"
        values[7] = values[7].replace("甲" * 88, f"{repeated}。{repeated}。" + "甲" * 45)
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "重复句或分句")

    def test_generic_filler_is_detected(self) -> None:
        values = valid_values()
        values[4] = "将作用机制、临床用途与风险提示联系起来进行辨析。"
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "通用填充语")

    def test_lecture_meta_language_is_detected(self) -> None:
        values = valid_values()
        values[7] = values[7].replace("甲" * 20, "围绕受体作用进行学习" + "甲" * 10, 1)
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "教学过程套话")

    def test_rigid_homework_filler_is_detected(self) -> None:
        values = valid_values()
        values[9] = values[9].replace("根据给定曲线", "围绕本课根据给定曲线")
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "僵化或笼统表达")

    def test_homework_labels_are_not_mandatory(self) -> None:
        values = valid_values()
        self.assertFalse(any("必须包含" in error for error in validate_loaded_workbook(make_workbook([values]), 1)))

    def test_copied_chapter_title_is_detected(self) -> None:
        values = valid_values()
        values[3] = "第一章：测试章节的重点知识。"
        self.assert_has(validate_loaded_workbook(make_workbook([values]), 1), "不得复述完整章节标题")

    def test_exact_duplicate_lessons_are_detected(self) -> None:
        first = valid_values()
        second = valid_values()
        second[0] = "3月5日"
        second[1] = "第二章：另一测试章节"
        errors = validate_loaded_workbook(make_workbook([first, second]), 2)
        self.assert_has(errors, "完全重复")

    def test_near_duplicate_template_is_detected(self) -> None:
        first = valid_values()
        second = valid_values()
        second[0] = "3月5日"
        second[1] = "第二章：另一测试章节"
        second[2] = first[2].replace("受体", "酶")
        errors = validate_loaded_workbook(make_workbook([first, second]), 2)
        self.assert_has(errors, "高相似课次")

    def test_repeated_homework_task_mix_is_detected(self) -> None:
        rows = []
        for index in range(5):
            values = valid_values()
            values[0] = f"3月{index + 1}日"
            values[1] = f"第{index + 1}章：不同章节"
            values[2] = values[2].replace("受体", f"受体{index}")
            values[3] = values[3].replace("曲线", f"曲线{index}")
            values[4] = values[4].replace("曲线", f"曲线{index}")
            values[5] = values[5].replace("受体", f"受体{index}")
            values[6] = chr(0x4e00 + index) * 100
            values[7] = values[7].replace("甲", chr(0x5000 + index))
            values[8] = values[8].replace("概念", f"概念{index}")
            values[9] = (
                f"1、完成教材中第{index + 1}组习题，并写出具体判断依据。\n"
                f"2、制作药物{index}与对照药的比较卡，注明机制和用途。\n"
                f"3、查阅药物{index}说明书，摘录三条用药注意事项。"
            )
            rows.append(values)
        self.assert_has(validate_loaded_workbook(make_workbook(rows), 5), "任务组合")


if __name__ == "__main__":
    unittest.main()
